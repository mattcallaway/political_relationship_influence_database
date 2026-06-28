import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os

os.makedirs('data/imports', exist_ok=True)
os.makedirs('scripts', exist_ok=True)

wb = openpyxl.Workbook()
wb.remove(wb.active) # Remove default sheet

def apply_header_style(ws, headers):
    ws.append(headers)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

# 1. Controlled Vocabularies Sheet
ws_vocab = wb.create_sheet(title="Controlled_Vocabularies")
vocab_headers = ["Vocabulary_Category", "Code", "Label", "Description", "Is_Active"]
apply_header_style(ws_vocab, vocab_headers)
vocab_data = [
    ["entity_type", "PERSON", "Person", "Individual human being", True],
    ["entity_type", "ORGANIZATION", "Organization", "Corporation, PAC, union, agency, or group", True],
    ["org_category", "CAMPAIGN_COMMITTEE", "Campaign Committee", "Registered candidate or ballot measure committee", True],
    ["org_category", "DEVELOPER_CORP", "Development Corporation", "Real estate development entity", True],
    ["org_category", "GOVT_AGENCY", "Government Agency", "County, city, or regional public body", True],
    ["org_category", "LOBBYING_FIRM", "Lobbying Firm", "Registered public affairs/lobbying firm", True],
    ["source_type", "FORM_460", "Form 460 Campaign Statement", "FPPC Campaign Disclosure Statement", True],
    ["source_type", "MEETING_MINUTES", "Meeting Minutes", "Official board or council meeting minutes", True],
    ["source_type", "STAFF_REPORT", "Board Staff Report", "County staff board recommendation packet", True],
    ["claim_type", "FACTUAL", "Factual Assertion", "Directly documented fact from public record", True],
    ["claim_type", "INTERPRETIVE", "Interpretive Assertion", "Analytical interpretation by researcher", True],
    ["claim_type", "ALLEGATION", "Allegation", "Formally reported claim under inquiry", True],
    ["predicate", "CONTRIBUTED_TO", "Contributed To", "Financial or non-monetary contribution", True],
    ["predicate", "EMPLOYED_BY", "Employed By", "Employment or contractual role", True],
    ["predicate", "REPRESENTED_CLIENT", "Represented Client", "Lobbying or legal representation", True],
    ["predicate", "VOTED_FOR", "Voted For", "Cast affirmative vote on motion or ordinance", True],
]
for row in vocab_data:
    ws_vocab.append(row)

# 2. Entities Sheet
ws_entities = wb.create_sheet(title="Entities")
apply_header_style(ws_entities, ["Legacy_ID", "Canonical_Name", "Entity_Type", "Status", "Notes"])
entities_data = [
    ["LEG_E_001", "Lynda Hopkins", "PERSON", "APPROVED", "Sonoma County Supervisor District 5"],
    ["LEG_E_002", "James Gore", "PERSON", "APPROVED", "Sonoma County Supervisor District 4"],
    ["LEG_E_003", "Re-Elect Lynda Hopkins for Supervisor 2024", "ORGANIZATION", "APPROVED", "FPPC ID 1424912"],
    ["LEG_E_004", "Sonoma Coast Development LLC", "ORGANIZATION", "APPROVED", "Commercial developer in West County"],
    ["LEG_E_005", "Pacific Public Affairs Group", "ORGANIZATION", "APPROVED", "Regional lobbying and advocacy firm"],
    ["LEG_E_006", "Sonoma County Board of Supervisors", "ORGANIZATION", "APPROVED", "County governing body"],
]
for row in entities_data:
    ws_entities.append(row)

# 3. Persons Sheet
ws_persons = wb.create_sheet(title="Persons")
apply_header_style(ws_persons, ["Legacy_ID", "Full_Name", "First_Name", "Last_Name", "Occupation", "Public_Role", "Jurisdiction"])
persons_data = [
    ["LEG_E_001", "Lynda Hopkins", "Lynda", "Hopkins", "County Supervisor", "Supervisor D5", "Sonoma County"],
    ["LEG_E_002", "James Gore", "James", "Gore", "County Supervisor", "Supervisor D4", "Sonoma County"],
]
for row in persons_data:
    ws_persons.append(row)

# 4. Organizations Sheet
ws_orgs = wb.create_sheet(title="Organizations")
apply_header_style(ws_orgs, ["Legacy_ID", "Legal_Name", "Common_Name", "Org_Category", "Committee_ID", "Jurisdiction"])
orgs_data = [
    ["LEG_E_003", "Re-Elect Lynda Hopkins for Supervisor 2024", "Hopkins for Supervisor", "CAMPAIGN_COMMITTEE", "1424912", "Sonoma County"],
    ["LEG_E_004", "Sonoma Coast Development LLC", "Sonoma Coast Dev", "DEVELOPER_CORP", "", "California"],
    ["LEG_E_005", "Pacific Public Affairs Group", "Pacific Public Affairs", "LOBBYING_FIRM", "LOB-7092", "Sonoma County"],
    ["LEG_E_006", "Sonoma County Board of Supervisors", "BOS", "GOVT_AGENCY", "", "Sonoma County"],
]
for row in orgs_data:
    ws_orgs.append(row)

# 5. Sources Sheet
ws_sources = wb.create_sheet(title="Sources")
apply_header_style(ws_sources, ["Legacy_Source_ID", "Title", "Source_Type", "Publisher", "Filing_Date", "URL", "Reliability"])
sources_data = [
    ["LEG_SRC_101", "Form 460 - Hopkins for Supervisor 2024 (Late 2023)", "FORM_460", "Sonoma County Registrar of Voters", "2024-01-31", "https://sccounty.gov/filings/460_1424912_2023.pdf", "OFFICIAL_PUBLIC_RECORD"],
    ["LEG_SRC_102", "Sonoma County BOS Meeting Minutes - Dec 12 2023", "MEETING_MINUTES", "Sonoma County Board of Supervisors", "2023-12-12", "https://sonoma.legistar.com/View.ashx?M=M&ID=1062948", "OFFICIAL_PUBLIC_RECORD"],
]
for row in sources_data:
    ws_sources.append(row)

# 6. Assertions Sheet
ws_assertions = wb.create_sheet(title="Assertions")
apply_header_style(ws_assertions, ["Legacy_Assertion_ID", "Subject_Legacy_ID", "Predicate", "Object_Legacy_ID", "Literal_Value", "Claim_Type", "Confidence", "Source_Legacy_ID", "Source_Page"])
assertions_data = [
    ["LEG_AST_501", "LEG_E_004", "CONTRIBUTED_TO", "LEG_E_003", "", "FACTUAL", "VERIFIED", "LEG_SRC_101", 3],
    ["LEG_AST_502", "LEG_E_001", "EMPLOYED_BY", "LEG_E_006", "", "FACTUAL", "VERIFIED", "LEG_SRC_102", 1],
]
for row in assertions_data:
    ws_assertions.append(row)

# 7. Contributions Sheet
ws_contribs = wb.create_sheet(title="Contributions")
apply_header_style(ws_contribs, ["Legacy_Trans_ID", "Filer_Committee_ID", "Donor_Legacy_ID", "Donor_Raw_Name", "Transaction_Date", "Amount", "Schedule", "Source_Legacy_ID", "Source_Page", "Review_Status"])
contribs_data = [
    ["LEG_TRN_001", "LEG_E_003", "LEG_E_004", "Sonoma Coast Development LLC", "2023-11-15", 2500.00, "Schedule A", "LEG_SRC_101", 3, "APPROVED"],
]
for row in contribs_data:
    ws_contribs.append(row)

# 8. Expenditures Sheet
ws_expenditures = wb.create_sheet(title="Expenditures")
apply_header_style(ws_expenditures, ["Legacy_Trans_ID", "Filer_Committee_ID", "Payee_Legacy_ID", "Payee_Raw_Name", "Transaction_Date", "Amount", "Description", "Schedule", "Source_Legacy_ID", "Source_Page", "Review_Status"])
exp_data = [
    ["LEG_TRN_002", "LEG_E_003", "LEG_E_005", "Pacific Public Affairs Group", "2023-12-01", 5000.00, "Campaign Strategy Consulting", "Schedule E", "LEG_SRC_101", 6, "APPROVED"],
]
for row in exp_data:
    ws_expenditures.append(row)

# 9. Research Workflow Sheet
ws_workflow = wb.create_sheet(title="Research_Workflow")
apply_header_style(ws_workflow, ["Legacy_Task_ID", "Task_Title", "Assigned_Researcher", "Priority", "Status", "Notes"])
workflow_data = [
    ["LEG_TSK_901", "Verify Form 460 Schedule A corporate donors", "M. Callaway", "HIGH", "IN_PROGRESS", "Cross-reference California Secretary of State corporate filings for LLC members"],
]
for row in workflow_data:
    ws_workflow.append(row)

# 10. Legacy ID Mappings Sheet
ws_legacy = wb.create_sheet(title="Legacy_ID_Mappings")
apply_header_style(ws_legacy, ["Legacy_ID", "Source_Sheet", "Target_Model", "System_Public_ID_Prefix"])
legacy_data = [
    ["LEG_E_001", "Entities", "Person / Entity", "P"],
    ["LEG_E_002", "Entities", "Person / Entity", "P"],
    ["LEG_E_003", "Entities", "Organization / Entity", "ORG"],
    ["LEG_E_004", "Entities", "Organization / Entity", "ORG"],
    ["LEG_E_005", "Entities", "Organization / Entity", "ORG"],
    ["LEG_E_006", "Entities", "Organization / Entity", "ORG"],
    ["LEG_SRC_101", "Sources", "Source", "SRC"],
    ["LEG_SRC_102", "Sources", "Source", "SRC"],
    ["LEG_AST_501", "Assertions", "Assertion", "AST"],
    ["LEG_AST_502", "Assertions", "Assertion", "AST"],
    ["LEG_TRN_001", "Contributions", "Contribution", "CON"],
    ["LEG_TRN_002", "Expenditures", "Expenditure", "EXP"],
]
for row in legacy_data:
    ws_legacy.append(row)

excel_path = "data/imports/Sonoma_Political_Influence_Database_v5_consolidated.xlsx"
wb.save(excel_path)
print(f"Workbook successfully generated at {excel_path}")
