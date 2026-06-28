import openpyxl
import json

wb_path = "data/imports/Sonoma_Political_Influence_Database_v5_consolidated.xlsx"
wb = openpyxl.load_workbook(wb_path, data_only=True)

audit_results = {}

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        continue
    headers = [str(cell) for cell in rows[0] if cell is not None]
    row_count = len(rows) - 1
    sample_rows = rows[1:4] if row_count > 0 else []
    
    audit_results[sheetname] = {
        "column_count": len(headers),
        "columns": headers,
        "total_records": row_count,
        "sample_data": sample_rows
    }

print(json.dumps(audit_results, indent=2, default=str))
